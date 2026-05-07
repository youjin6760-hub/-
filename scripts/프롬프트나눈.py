import base64
import json
import mimetypes
import os
import re
import unicodedata

from pathlib import Path
import time

from openai import OpenAI, RateLimitError
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "raw"

REVIEWED_DIR = ROOT / "reviewed_json"
REVIEWED_DIR.mkdir(exist_ok=True)

# ISSUE_XLSX_PATH = ROOT / "json_issue_index.xlsx"
ISSUE_XLSX_PATH = ROOT / "sqld10000000000.xlsx"
FORMULA_XLSX_PATH = ROOT / "sqld_ex10000000000.xlsx"

MODEL_NAME = "gpt-5.4-mini"
MAX_TOKENS = 2500

MAX_QUESTION_IMAGES = 2
MAX_CHOICE_IMAGES = 4
MAX_EXPLANATION_IMAGES = 2


def load_prompt_content():
    with open(ROOT / "내용검수프롬프트.txt", "r", encoding="utf-8") as f:
        return f.read()


def load_prompt_format():
    with open(ROOT / "형식검수프롬프트.txt", "r", encoding="utf-8") as f:
        return f.read()
    

def guess_media_type(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime in {"image/png", "image/jpeg", "image/gif", "image/webp"}:
        return mime
    return "image/png"


def make_text_block(text: str) -> dict:
    return {
        "type": "input_text",
        "text": text
    }


def make_image_block(path_str: str, label: str) -> list[dict]:
    path = Path(path_str)

    if not path.exists():
        return [make_text_block(f"[이미지 파일 누락] {label}: {path_str}")]

    media_type = guess_media_type(path)

    with open(path, "rb") as f:
        encoded = base64.b64encode(f.read()).decode("utf-8")

    return [
        make_text_block(f"[첨부 이미지] {label}: {path.name}"),
        {
            "type": "input_image",
            "image_url": f"data:{media_type};base64,{encoded}"
        }
    ]

def safe_sheet_name(name: str) -> str:
    if not name:
        return "issues"

    name = str(name)

    for ch in ['\\', '/', '*', '?', ':', '[', ']']:
        name = name.replace(ch, "_")

    return name[:31]


def safe_excel_text(v):
    if v is None:
        return ""

    v = str(v)
    v = unicodedata.normalize("NFC", v)

    # 엑셀에서 문제 되는 제어문자만 제거
    v = re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]", "", v)

    # 줄바꿈은 셀 깨짐 방지용으로 공백 처리
    v = v.replace("\r", " ").replace("\n", " ")

    return v.strip()


ISSUE_HEADERS = [
    "question_id",
    "file_name",
    "question_no",
    "subject_name",
    "sub_title",
    "severity",
    "issue_area",
    "issue_type",
    "reason",
    "suggestion",
    "confidence"
]

FORMULA_HEADERS = [
    "question_id",
    "file_name",
    "question_no",
    "subject_name",
    "sub_title",
    "has_formula_explanation",
    "explanation_image_count",
    "capture_mode",
    "needs_manual_review",
    "explanation_images",
    "formula_reason",
]

def last_non_empty_row(ws):
    for row in range(ws.max_row, 0, -1):
        if any(
            ws.cell(row=row, column=col).value not in (None, "")
            for col in range(1, ws.max_column + 1)
        ):
            return row
    return 0


def save_issue_rows_to_xlsx(issue_rows):
    if not issue_rows:
        print("[ISSUE XLSX] 저장할 오류 없음")
        return

    is_existing_file = ISSUE_XLSX_PATH.exists()

    if is_existing_file:
        wb = load_workbook(ISSUE_XLSX_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws_map = {ws.title: ws for ws in wb.worksheets}

    # 파일 열었을 때 이미 존재하던 시트만 기록
    existing_sheet_names = set(ws_map.keys())

    # 이번 실행에서 기존 데이터와 신규 데이터 사이 공백을 이미 넣은 시트
    separated_sheets = set()

    for row_data in issue_rows:
        sheet_base = row_data.get("sub_title") or row_data.get("subject_name") or "issues"
        sheet_name = safe_sheet_name(sheet_base)

        if sheet_name not in ws_map:
            ws = wb.create_sheet(sheet_name)
            ws_map[sheet_name] = ws

            for col_idx, header in enumerate(ISSUE_HEADERS, start=1):
                ws.cell(row=1, column=col_idx, value=header)

            next_row = 2

        else:
            ws = ws_map[sheet_name]
            last_row = last_non_empty_row(ws)

            # 기존 파일에 원래 있던 시트이고, 기존 데이터가 있을 때만 한 줄 띄움
            if (
                sheet_name in existing_sheet_names
                and last_row > 1
                and sheet_name not in separated_sheets
            ):
                next_row = last_row + 2
                separated_sheets.add(sheet_name)
            else:
                next_row = last_row + 1

        row_values = [
            row_data["question_id"],
            row_data["file_name"],
            row_data["question_no"],
            row_data["subject_name"],
            row_data["sub_title"],
            row_data["severity"],
            row_data["issue_area"],
            row_data["issue_type"],
            row_data["reason"],
            row_data["suggestion"],
            row_data["confidence"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=next_row, column=col_idx, value=safe_excel_text(value))

    wb.save(ISSUE_XLSX_PATH)
    print(f"[ISSUE XLSX 저장 완료] {ISSUE_XLSX_PATH}")
              
              
def has_formula_explanation(data: dict):
    explanation_images = data.get("explanation_images", []) or []
    meta = data.get("explanation_capture_meta", {}) or {}

    reasons = []

    if explanation_images:
        reasons.append("해설 스크린샷 존재")

    capture_mode = meta.get("capture_mode")

    if meta.get("attempted") is True and capture_mode not in ("not_needed", "not_attempted"):
        reasons.append(f"collect 수식 조건 감지: capture_mode={capture_mode}")

    return bool(reasons), ", ".join(reasons)


def save_formula_rows_to_xlsx(formula_rows):
    if not formula_rows:
        print("[FORMULA XLSX] 저장할 수식 해설 문제 없음")
        return

    is_existing_file = FORMULA_XLSX_PATH.exists()

    if is_existing_file:
        wb = load_workbook(FORMULA_XLSX_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws_map = {ws.title: ws for ws in wb.worksheets}
    existing_sheet_names = set(ws_map.keys())
    separated_sheets = set()

    for row_data in formula_rows:
        sheet_base = row_data.get("sub_title") or row_data.get("subject_name") or "formula"
        sheet_name = safe_sheet_name(sheet_base)

        if sheet_name not in ws_map:
            ws = wb.create_sheet(sheet_name)
            ws_map[sheet_name] = ws

            for col_idx, header in enumerate(FORMULA_HEADERS, start=1):
                ws.cell(row=1, column=col_idx, value=header)

            next_row = 2
        else:
            ws = ws_map[sheet_name]
            last_row = last_non_empty_row(ws)

            if (
                sheet_name in existing_sheet_names
                and last_row > 1
                and sheet_name not in separated_sheets
            ):
                next_row = last_row + 2
                separated_sheets.add(sheet_name)
            else:
                next_row = last_row + 1

        row_values = [
            row_data["question_id"],
            row_data["file_name"],
            row_data["question_no"],
            row_data["subject_name"],
            row_data["sub_title"],
            row_data["has_formula_explanation"],
            row_data["explanation_image_count"],
            row_data["capture_mode"],
            row_data["needs_manual_review"],
            row_data["explanation_images"],
            row_data["formula_reason"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=next_row, column=col_idx, value=safe_excel_text(value))

    wb.save(FORMULA_XLSX_PATH)
    print(f"[FORMULA XLSX 저장 완료] {FORMULA_XLSX_PATH}")
    
              
# =========================
# API 입력 구성
# =========================
def build_user_content(prompt: str, question_data: dict, mode: str) -> list[dict]:
    blocks: list[dict] = []

    data = question_data.get("data", {}) or {}
    image_elements = data.get("image_elements", []) or []
    explanation_images = data.get("explanation_images", []) or []

    has_question_image = bool(data.get("has_question_image"))
    has_choice_image = bool(data.get("has_choice_image"))
    has_image = bool(data.get("has_image"))
    body = data.get("body", "")
    extra_text = data.get("extra_text", "")
    choices = data.get("choices", [])

    choice_mapping_text = "\n".join(
        f"{idx}. {choice}"
        for idx, choice in enumerate(choices, start=1)
    )

    review_json = json.dumps({
        "body": data.get("body"),
        "extra_text": data.get("extra_text"),
        "choices": data.get("choices"),
        "answer": data.get("answer"),
        "explanation": data.get("explanation"),
        "explanation_capture_meta": data.get("explanation_capture_meta"),
        "image_elements": data.get("image_elements"),
        "explanation_images": data.get("explanation_images"),
    }, ensure_ascii=False, indent=2)
        
    if mode == "content":
        matching_instruction = (
            "\n[선지-해설 강제 매칭 검수]\n"
            "다음 choices와 explanation의 선지별 해설을 반드시 번호별로 1:1 비교하세요.\n"
            "번호가 같더라도 설명 내용이 다른 선지를 설명하면 '선지-해설 불일치'입니다.\n\n"
            "이 검수는 반드시 수행해야 하며, 하나라도 불일치가 있으면 반드시 오류로 기록해야 합니다.\n\n"
            f"[choices]\n{choice_mapping_text}\n\n"
            f"[explanation]\n{data.get('explanation', '')}\n\n"
        )
    else:
        matching_instruction = (
            "\n[형식 검수용 선지 정보]\n"
            "아래 choices와 explanation은 선지 해설의 형식 확인용으로만 사용하세요.\n"
            "선지 내용의 정오나 해설 논리 불일치는 판단하지 마세요.\n\n"
            f"[choices]\n{choice_mapping_text}\n\n"
            f"[explanation]\n{data.get('explanation', '')}\n\n"
        )
        
    question_items = [x for x in image_elements if x.get("location") == "question"]
    choice_items = [x for x in image_elements if x.get("location") == "choice"]

    if mode == "content":
        extra_instruction = (
            "- 텍스트 문제라도 반드시 문제 성립/정답/해설 논리 검수를 수행하세요.\n"
            "- 형식 오류보다 문제 자체 오류를 우선 검수하세요.\n"
            "- 이미지가 문제 풀이에 필수인데 없거나 잘못된 경우 반드시 문제 성립 오류로 판단하세요.\n"
            "- 문제 이미지, 선지 이미지, 해설 스크린샷은 서로 독립적으로 검수하세요.\n"
            "- 문제 풀이가 불가능하면 다른 판단 없이 즉시 문제 성립 오류로 기록하세요.\n"
        )
    elif mode == "format":
        extra_instruction = (
            "- 내용 정오, 정답 정합성, 문제 성립 여부는 검수하지 마세요.\n"
            "- 해설 시작 문장, 선지별 해설 구조, 결론 문장, 정답 문장 중복, 긴 해설 수동 검토만 판단하세요.\n"
            "- 이미지는 내용 검수에 사용하지 말고, 해설 구조 확인 참고용으로만 사용하세요.\n"
        )
    else:
        extra_instruction = ""
        
    blocks.append(
        make_text_block(
            f"{prompt}\n\n"
            f"[문제 데이터 JSON]\n"
            f"{review_json}\n\n"
            f"{matching_instruction}"
            f"[검수용 추가 정보]\n"
            f"{extra_instruction}"
            f"- has_image: {has_image}\n"
            f"- has_question_image: {has_question_image}\n"
            f"- has_choice_image: {has_choice_image}\n"
            f"- question_image_count: {len(question_items)}\n"
            f"- choice_image_count: {len(choice_items)}\n"
            f"- explanation_image_count: {len(explanation_images)}\n"
            f"- body: {body}\n"
            f"- extra_text: {extra_text}\n"
            f"- choices_count: {len(choices)}\n"
            f"- 설명 문장, 마크다운, 코드블록 없이 순수 JSON만 반환하세요.\n"
        )
    )

    if question_items:
        blocks.append(make_text_block("[문제 이미지 검수 대상]"))
        for idx, item in enumerate(question_items[:MAX_QUESTION_IMAGES], start=1):
            saved_path = item.get("saved_path", "")
            near = item.get("caption_or_near_text", "") or item.get("ocr_or_extracted_text", "")
            label = f"문제 이미지 {idx}"
            if near:
                label += f" / {near}"
            blocks.extend(make_image_block(saved_path, label))
    else:
        blocks.append(make_text_block("[문제 이미지 없음]"))

    if choice_items:
        blocks.append(make_text_block("[선지 이미지 검수 대상]"))
        for idx, item in enumerate(choice_items[:MAX_CHOICE_IMAGES], start=1):
            saved_path = item.get("saved_path", "")
            near = item.get("caption_or_near_text", "") or item.get("ocr_or_extracted_text", "")
            label = f"선지 이미지 {idx}"
            if near:
                label += f" / {near}"
            blocks.extend(make_image_block(saved_path, label))
    else:
        blocks.append(make_text_block("[선지 이미지 없음]"))

    if explanation_images:
        blocks.append(make_text_block("[해설 스크린샷 검수 대상]"))
        for idx, path_str in enumerate(explanation_images[:MAX_EXPLANATION_IMAGES], start=1):
            blocks.extend(make_image_block(str(path_str), f"해설 스크린샷 {idx}"))
    else:
        blocks.append(make_text_block("[해설 스크린샷 없음]"))

    return blocks


def parse_review_response(text: str, question_id: str) -> dict:
    try:
        text = text.strip()

        if text.startswith("```"):
            text = re.sub(r"^```(?:json|JSON)?\s*", "", text)
            text = re.sub(r"\s*```$", "", text)

        return json.loads(text)

    except Exception:
        return {
            "question_id": question_id,
            "content_issues": [{
                "type": "기타 내용 오류",
                "reason": "검수 결과를 JSON으로 파싱하지 못했습니다.",
                "suggestion": "검수 프롬프트 출력 형식을 재확인하세요.",
                "confidence": 1.0
            }],
            "format_issues": [],
            "summary": {
                "has_issue": True,
                "severity": "medium"
            }
        }


def merge_severity(sev1: str, sev2: str) -> str:
    order = {
        "low": 1,
        "medium": 2,
        "high": 3
    }

    s1 = sev1 if sev1 in order else "low"
    s2 = sev2 if sev2 in order else "low"

    return s1 if order[s1] >= order[s2] else s2


def is_false_positive_normal_issue(issue: dict) -> bool:
    text = (
        str(issue.get("type", "")) + " " +
        str(issue.get("reason", "")) + " " +
        str(issue.get("suggestion", ""))
    )

    normal_phrases = [
        "오류는 없습니다",
        "형식 오류는 없습니다",
        "내용 오류는 없습니다",
        "문제는 없습니다",
        "이상 없습니다",
        "정상입니다",
        "정상입니다.",
        "일치합니다",
        "일치하므로",
        "불일치는 없습니다",
        "불일치가 없습니다",
        "해당 없음",
        "해당 없음으로 보입니다",
        "수정할 필요 없습니다",
        "수정할 필요는 없습니다",
        "유지해도 됩니다",
        "문제 없어 보입니다",
        "문제는 없어 보입니다",
        "형식상 문제는 없어 보입니다",
        "확인 가능한 형식 오류는 없습니다",
        "확인 가능한 오류는 없습니다",
        "오류로 판단하지 않습니다",
        "오류로 보지 않습니다",
    ]

    return any(p in text for p in normal_phrases)


def normalize_issue_key(issue: dict) -> str:
    issue_type = str(issue.get("type", ""))
    reason = str(issue.get("reason", ""))
    suggestion = str(issue.get("suggestion", ""))

    text = f"{issue_type} {reason} {suggestion}"

    # 해설 시작 형식 오류 묶기
    if (
        issue_type == "해설 시작 형식 오류"
        or "해설 시작" in text
        or "explanation 시작" in text
        or "첫 문장" in text
        or "답은" in text
        or "정답은 X번입니다" in text
        or "정답은 [숫자]번입니다" in text
    ):
        return "해설 시작 형식 오류"

    # 정답 문장 중복 묶기
    if (
        issue_type == "정답 문장 중복"
        or (
            "정답은" in text
            and any(w in text for w in ["반복", "중복", "다시 등장", "한 번 더"])
        )
    ):
        return "정답 문장 중복"

    # 선지 해설 형식 오류 묶기
    if (
        issue_type == "선지 해설 형식 오류"
        or "선지 해설" in text
        or "X. 선지 전체 문장" in text
    ):
        return "선지 해설 형식 오류"

    # 문체 오류 묶기
    if any(w in text for w in ["문체", "존댓말", "반말", "구어체", "평서체"]):
        return "문체 오류"

    # 정상 판단으로 작성된 issue 제거용 key
    if any(w in text for w in [
        "오류는 없습니다",
        "형식 오류는 없습니다",
        "내용 오류는 없습니다",
        "문제는 없습니다",
        "이상 없습니다",
        "정상입니다",
        "일치합니다",
        "불일치는 없습니다",
        "불일치가 없습니다",
        "해당 없음",
        "해당 없음으로 보입니다",
        "수정할 필요 없습니다",
        "수정할 필요는 없습니다",
        "유지해도 됩니다",
        "문제 없어 보입니다",
        "문제는 없어 보입니다",
        "형식상 문제는 없어 보입니다",
        "확인 가능한 형식 오류는 없습니다",
        "확인 가능한 오류는 없습니다",
    ]):
        return "정상 판단"

    return f"{issue_type}|{reason[:50]}"


def dedupe_issues(issues: list) -> list:
    if not issues:
        return []

    result = []
    seen = set()

    for issue in issues:
        key = normalize_issue_key(issue)

        # 정상 판단은 오류로 저장하지 않음
        if key == "정상 판단":
            continue

        if key in seen:
            continue

        seen.add(key)
        result.append(issue)

    return result


def merge_review_results(question_data: dict, content_result: dict, format_result: dict) -> dict:
    content_issues = content_result.get("content_issues", []) or []
    format_issues = format_result.get("format_issues", []) or []

    # 1) 정상 판단을 issue로 잘못 작성한 항목 먼저 제거
    content_issues = [
        issue for issue in content_issues
        if not is_false_positive_normal_issue(issue)
    ]

    format_issues = [
        issue for issue in format_issues
        if not is_false_positive_normal_issue(issue)
    ]

    # 2) 같은 원인의 오류 중복 제거
    content_issues = dedupe_issues(content_issues)
    format_issues = dedupe_issues(format_issues)

    has_issue = bool(content_issues or format_issues)

    severity = merge_severity(
        content_result.get("summary", {}).get("severity", "low"),
        format_result.get("summary", {}).get("severity", "low")
    )

    if not has_issue:
        severity = "low"

    return {
        **question_data,
        "content_issues": content_issues,
        "format_issues": format_issues,
        "summary": {
            "has_issue": has_issue,
            "severity": severity
        }
    }
    
    
def add_duplicate_answer_sentence_issue(reviewed: dict, question_data: dict) -> dict:
    explanation = question_data.get("data", {}).get("explanation", "") or ""

    pattern = r"정답은\s*[1-5]\s*번입니다\.?"
    matches = list(re.finditer(pattern, explanation))

    if not matches:
        return reviewed

    # 공백 제거 후 첫 문장이 정답 문장으로 시작하는지 확인
    stripped = explanation.lstrip()
    first_match = matches[0]

    starts_correctly = first_match.start() == len(explanation) - len(stripped)

    # 오류 조건:
    # 1) 정답 문장이 2개 이상 있음
    # 2) 정답 문장이 1개만 있어도 첫 문장 위치가 아님
    if len(matches) == 1 and starts_correctly:
        return reviewed

    issue = {
        "type": "정답 문장 중복",
        "reason": "해설 첫 문장이 아닌 위치에 '정답은 X번입니다.' 형식의 문장이 등장하여 해설 형식 기준에 맞지 않습니다.",
        "suggestion": "해설 첫 문장을 정확히 '정답은 X번입니다.' 형식으로 작성하고, 해설 중간 또는 결론 문단 직전의 정답 문장은 삭제하세요.",
        "confidence": 0.95
    }

    reviewed.setdefault("format_issues", [])

    already_exists = any(
        normalize_issue_key(i) == "정답 문장 중복"
        for i in reviewed.get("format_issues", [])
    )

    if not already_exists:
        reviewed["format_issues"].append(issue)

    # 혹시 후처리 후에도 중복이 생겼을 경우 한 번 더 정리
    reviewed["format_issues"] = dedupe_issues(reviewed.get("format_issues", []))

    reviewed.setdefault("summary", {})
    reviewed["summary"]["has_issue"] = bool(
        reviewed.get("content_issues", []) or reviewed.get("format_issues", [])
    )

    if reviewed["summary"]["has_issue"] and reviewed["summary"].get("severity", "low") == "low":
        reviewed["summary"]["severity"] = "medium"

    return reviewed


def review_one(client, prompt: str, question_data: dict, mode: str) -> dict:
    content = build_user_content(prompt, question_data, mode)

    for attempt in range(5):
        try:
            resp = client.responses.create(
                model=MODEL_NAME,
                input=[
                    {
                        "role": "user",
                        "content": content
                    }
                ],
                temperature=0,
                max_output_tokens=MAX_TOKENS
            )
            break

        except RateLimitError:
            wait = 5 * (attempt + 1)
            print(f"[RateLimit] {wait}초 대기 후 재시도")
            time.sleep(wait)

        except Exception as e:
            print(f"[API 오류] {e}")
            raise

    else:
        raise RuntimeError("RateLimit 재시도 실패")

    text = resp.output_text.strip()
    return parse_review_response(text, question_data.get("question_id", ""))


def question_no_from_filename(path: Path):
    try:
        return int(path.stem.split("_")[-1])
    except Exception:
        return 999999


def main():
    load_dotenv(ROOT / ".env")
    
    api_key = os.getenv("CHATGPT_API")
    if not api_key:
        raise RuntimeError(".env에 CHATGPT_API가 필요합니다.")

    client = OpenAI(
        api_key=api_key,
        timeout=120
    )
    prompt_content = load_prompt_content()
    prompt_format = load_prompt_format()

    issue_rows = []
    formula_rows = []
    
    # # =========================
    # # 전체 JSON 검수
    # # =========================
    # raw_files = sorted(
    #     RAW_DIR.glob("*.json"),
    #     key=question_no_from_filename
    # )


    # =========================
    # 특정 회차만 검수
    # =========================
    TARGET_PREFIXES = [
    # "SQLD_61회_끝장_패키지_PT_2026_오답킬러_100제_데이터_모델링의_이해_SQLD_오답킬러_1과목_",
    # "SQLD_61회_끝장_패키지_PT_2026_오답킬러_100제_SQL_기본_및_활용_SQLD_오답킬러_2과목_1장_",
    # "SQLD_61회_끝장_패키지_PT_2026_오답킬러_100제_SQL_기본_및_활용_SQLD_오답킬러_2과목_2장_",
    # "SQLD_61회_끝장_패키지_PT_2026_오답킬러_100제_SQL_기본_및_활용_SQLD_오답킬러_2과목_3장_",
    # "SQLD_61회_끝장_패키지_PT_2026_최빈_완전정복_100제_데이터_모델링의_이해_SQLD_빈출_1과목_",
    # "SQLD_61회_끝장_패키지_PT_2026_최빈_완전정복_100제_SQL_기본_및_활용_SQLD_빈출_2과목_1장_",
    # "SQLD_61회_끝장_패키지_PT_2026_최빈_완전정복_100제_SQL_기본_및_활용_SQLD_빈출_2과목_2장_",
    # "SQLD_61회_끝장_패키지_PT_2026_최빈_완전정복_100제_SQL_기본_및_활용_SQLD_빈출_2과목_3장_",
    # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_26년_56회_기출복원문제_",
    # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_26년_57회_기출복원문제_",
    # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_26년_58회_기출복원문제_",
    # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_26년_59회_기출복원문제_",
    # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_26년_60회_기출복원문제_",
    "SQLD_61회_끝장_패키지_PT_2026_핵심개념_문제_SQL_기본_및_활용_SQLD_핵심개념체크_2과목_1장_",
    "SQLD_61회_끝장_패키지_PT_2026_핵심개념_문제_SQL_기본_및_활용_SQLD_핵심개념체크_2과목_2장_",
    "SQLD_61회_끝장_패키지_PT_2026_핵심개념_문제_SQL_기본_및_활용_SQLD_핵심개념체크_2과목_3장_",
    # "SQLD_예비문제수록_예상문제_데이터_모델링의_이해_SQLD_예비문제_1과목_"
    # "SQLD_예비문제수록_예상문제_SQL_기본_및_활용_SQLD_예비문제_2과목_"
    ]

    raw_files = [
        p for p in RAW_DIR.glob("*.json")
        if any(p.name.startswith(prefix) for prefix in TARGET_PREFIXES)
    ]

    raw_files = sorted(
        raw_files,
        key=lambda p: (
            next(
                (idx for idx, prefix in enumerate(TARGET_PREFIXES)
                if p.name.startswith(prefix)),
                999
            ),
            question_no_from_filename(p)
        )
    )
    
    
    # # =========================
    # # 지정 문제만 검수
    # # =========================
    # TARGET_QUESTIONS = {
    #     # "SQLD_61회_끝장_패키지_PT_2026_핵심개념_문제_SQL_기본_및_활용_SQLD_핵심개념체크_2과목_2장_": [41, 42],
    #     # "SQLD_61회_끝장_패키지_PT_2026_오답킬러_100제_SQL_기본_및_활용_SQLD_오답킬러_2과목_1장_": [14, 16],
    #     # "SQLD_61회_끝장_패키지_PT_2026_최신기출_문제_SQLD_25년_57회_기출복원문제": [28]
    #     "SQLD_예비문제수록_예상문제_데이터_모델링의_이해_SQLD_예비문제_1과목_": [34],
    #     # "SQLD_예비문제수록_예상문제_SQL_기본_및_활용_SQLD_예비문제_2과목_": [75]
    # }

    # raw_files = [
    #     p
    #     for p in RAW_DIR.glob("*.json")
    #     if any(
    #         p.name.startswith(prefix)
    #         and question_no_from_filename(p) in question_nos
    #         for prefix, question_nos in TARGET_QUESTIONS.items()
    #     )
    # ]

    # raw_files = sorted(
    #     raw_files,
    #     key=lambda p: (
    #         next(
    #             (idx for idx, prefix in enumerate(TARGET_QUESTIONS.keys())
    #             if p.name.startswith(prefix)),
    #             999
    #         ),
    #         question_no_from_filename(p)
    #     )
    # )
    
    #

    if not raw_files:
        print("[정보] raw 폴더에 검수할 JSON이 없습니다.")
        return

    for raw_file in raw_files:
        with open(raw_file, "r", encoding="utf-8") as f:
            question_data = json.load(f)
            

        try:
            content_result = review_one(
                client,
                prompt_content,
                question_data,
                mode="content"
            )

            format_result = review_one(
                client,
                prompt_format,
                question_data,
                mode="format"
            )

            merged = merge_review_results(
                question_data,
                content_result,
                format_result
            )

            merged = add_duplicate_answer_sentence_issue(merged, question_data)

        except Exception as e:
            print(f"[SKIP API 오류] {raw_file.name}: {e}")
            continue

        out_path = REVIEWED_DIR / raw_file.name

        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(merged, f, ensure_ascii=False, indent=2)
            
        # =========================
        # 오류 XLSX 기록 추가
        # =========================
        if merged["summary"].get("has_issue") is True:
            issues = []

            for issue in merged.get("content_issues", []):
                issues.append(("content", issue))

            for issue in merged.get("format_issues", []):
                issues.append(("format", issue))

            for issue_area, issue in issues:
                row_data = {
                    "question_id": merged.get("question_id", ""),
                    "file_name": raw_file.name,
                    "question_no": merged.get("question_no", ""),
                    "subject_name": merged.get("subject_name", ""),
                    "sub_title": merged.get("sub_title", ""),
                    "severity": merged.get("summary", {}).get("severity", ""),
                    "issue_area": issue_area,
                    "issue_type": issue.get("type", ""),
                    "reason": issue.get("reason", ""),
                    "suggestion": issue.get("suggestion", ""),
                    "confidence": issue.get("confidence", "")
                }

                print(f"[ISSUE 행 추가] {raw_file.name} / {issue_area} / {issue.get('type', '')}")

                issue_rows.append(row_data)

            print(f"[ISSUE XLSX 기록] {raw_file.name}")

        data = question_data.get("data", {}) or {}
        has_formula, formula_reason = has_formula_explanation(data)

        if has_formula:
            expl_images = data.get("explanation_images", []) or []
            expl_meta = data.get("explanation_capture_meta", {}) or {}

            formula_rows.append({
                "question_id": merged.get("question_id", ""),
                "file_name": raw_file.name,
                "question_no": merged.get("question_no", ""),
                "subject_name": merged.get("subject_name", ""),
                "sub_title": merged.get("sub_title", ""),
                "has_formula_explanation": True,
                "explanation_image_count": len(expl_images),
                "capture_mode": expl_meta.get("capture_mode", ""),
                "needs_manual_review": expl_meta.get("needs_manual_review", ""),
                "explanation_images": " | ".join(map(str, expl_images)),
                "formula_reason": formula_reason,
            })

            print(f"[FORMULA 행 추가] {raw_file.name} / {formula_reason}")

        print(
            f"[REVIEW 저장] {out_path.name} "
            f"/ issue={merged['summary'].get('has_issue')} "
            f"/ severity={merged['summary'].get('severity')}"
        )

        time.sleep(0.5)
    save_issue_rows_to_xlsx(issue_rows)
    save_formula_rows_to_xlsx(formula_rows)

if __name__ == "__main__":
    main()